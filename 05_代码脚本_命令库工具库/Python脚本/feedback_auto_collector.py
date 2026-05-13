#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
规则甄查 · 甄先生 v2.0 — 反馈自动采集器
=========================================
Playwright 模拟登录后台 → 抓取视频播放数据 → 导出 feedback_latest.xlsx
→ 自动触发 feedback_listener.py 自进化逻辑。

用法：
  python feedback_auto_collector.py --mock     # 生成模拟数据
  python feedback_auto_collector.py            # 真实采集（需配置 config.yaml）
"""

import json
import sys
import time
import os
from datetime import datetime
from pathlib import Path
from typing import Optional

import yaml
from playwright.sync_api import sync_playwright

_DATA_DIR = Path(__file__).resolve().parent.parent.parent / "data"
_CFG_PATH = Path(__file__).resolve().parent.parent.parent / "config.yaml"
_FEEDBACK_PATH = _DATA_DIR / "feedback_latest.xlsx"
_VERIFY_FLAG = _DATA_DIR / "scan_verified.flag"

if sys.platform == "win32" and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

# ── 反馈 Excel 表头（与 feedback_listener.py 一致） ──
_COLUMNS = [
    "case_id", "original_text", "refined_text", "strategy",
    "plays", "likes", "comments", "shares", "engagement_rate",
    "baseline_engagement",
]


# ──────────────────────────────────────────────
# 模拟数据生成
# ──────────────────────────────────────────────

def _generate_mock() -> list:
    """生成真实感模拟数据用于测试"""
    return [
        ("M001", "这是全网第一的赚钱秘籍，绝对不封号",
         "业内深耕的资产增值逻辑，长期合规运营",
         "standard", 5200, 89, 12, 5, 0.035),
        ("M002", "这是全网第一的赚钱秘籍，绝对不封号",
         "本内容为业内深耕的深度分析，致力于资产增值逻辑",
         "aggressive", 4800, 42, 8, 3, 0.035),
        ("M003", "零成本创业，一学就会的赚钱方法",
         "零成本起步，一学即会的资产增值逻辑",
         "standard", 8300, 215, 34, 28, 0.028),
        ("M004", "零成本创业，一学就会的赚钱方法",
         "零成本起步新路径，掌握资产增值核心逻辑",
         "aggressive", 7900, 98, 15, 11, 0.028),
        ("M005", "点击领取私信我，最快入账方式，绝对赚钱",
         "后台留言获取深度交流，核心入账路径",
         "standard", 3500, 155, 22, 18, 0.048),
        ("M006", "点击领取私信我，最快入账方式，绝对赚钱",
         "欢迎后台交流，获取核心入账路径与资产增值方案",
         "aggressive", 3100, 62, 9, 6, 0.048),
        # 对照：低风险文案，互动率正常
        ("M007", "今天教大家一个实用小技巧，学会能省不少钱",
         "今天分享一个实用小技巧，掌握后可优化日常开支",
         "standard", 12000, 580, 92, 45, 0.052),
        ("M008", "今天教大家一个实用小技巧，学会能省不少钱",
         "今天分享一个实用技巧，帮助你优化日常开支结构",
         "aggressive", 11500, 510, 78, 40, 0.052),
    ]


def _mock_feedback():
    """写入模拟数据并触发自进化"""
    print("=" * 48)
    print("  规则甄查 · 甄先生 v2.0 — 反馈自动采集 [模拟]")
    print("=" * 48)

    data = _generate_mock()
    _write_excel(data)

    # 触发自进化
    _trigger_evolution()


# ──────────────────────────────────────────────
# Playwright 真实采集
# ──────────────────────────────────────────────

def _real_collect(cfg: dict) -> list:
    """
    Playwright 登录后台并采集数据。
    需在 config.yaml 配置 login_url / data_url / username / password。
    """
    fc = cfg.get("feedback_collector", {})
    login_url = fc.get("login_url", "")
    data_url = fc.get("data_url", "")
    username = fc.get("username", "")
    password = fc.get("password", "")

    if not login_url:
        print("[FAIL] config.yaml feedback_collector.login_url 未配置")
        return []

    print(f"\n[PLAYWRIGHT] 登录 {login_url}")

    data = []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True, args=["--no-sandbox"])
        ctx = browser.new_context(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 Chrome/124.0.0.0 Safari/537.36"
            ),
            locale="zh-CN",
        )
        page = ctx.new_page()

        try:
            # 登录
            page.goto(login_url, wait_until="domcontentloaded", timeout=30000)
            page.wait_for_timeout(2000)

            # QR 验证挂起检测（先检查是否需要扫码）
            if _is_qr_page(page):
                _wait_qr_scan(page)

            # 填充凭据（通用模式，可被具体页面覆盖）
            if username and password:
                page.fill("input[name='username'], input[name='email'], input[type='text']",
                          username)
                page.fill("input[name='password'], input[type='password']", password)
                page.click("button[type='submit'], .login-btn, .submit-btn")

            page.wait_for_timeout(3000)

            # 提交后可能再次出现二维码（二次验证）
            if _is_qr_page(page):
                _wait_qr_scan(page)

            # 导航到数据页
            if data_url:
                page.goto(data_url, wait_until="domcontentloaded", timeout=30000)
                page.wait_for_timeout(3000)

            # 提取表格数据（通用模式）
            rows = page.query_selector_all("table tbody tr, .data-table tr, .el-table__body-wrapper tr")
            print(f"  [TABLE] 发现 {len(rows)} 行数据")

            for i, row in enumerate(rows[:50]):
                cells = row.query_selector_all("td, .el-table__cell")
                vals = [c.inner_text().strip() for c in cells]
                if len(vals) >= 5:
                    data.append({
                        "case_id": f"R{i+1:03d}",
                        "original_text": vals[0] if len(vals) > 0 else "",
                        "refined_text": vals[1] if len(vals) > 1 else "",
                        "strategy": vals[2] if len(vals) > 2 else "standard",
                        "plays": _parse_num(vals[3]) if len(vals) > 3 else 0,
                        "likes": _parse_num(vals[4]) if len(vals) > 4 else 0,
                        "comments": _parse_num(vals[5]) if len(vals) > 5 else 0,
                        "shares": _parse_num(vals[6]) if len(vals) > 6 else 0,
                    })

        except Exception as e:
            print(f"  [FAIL] 采集异常: {e}")
        finally:
            ctx.close()
            browser.close()

    if not data:
        print("[WARN] 未采集到数据，降级使用模拟数据")
        data = _generate_mock()

    return data


# ──────────────────────────────────────────────
# QR 验证挂起机制
# ──────────────────────────────────────────────

def _is_qr_page(page) -> bool:
    """检测当前页面是否为二维码登录页"""
    indicators = [
        "img[src*='qrcode'], img[src*='qr_code'], img[src*='QR']",
        ".qrcode, .qr-code, .qr_code",
        "canvas[class*='qr']",
        "[data-role='qrcode']",
    ]
    for sel in indicators:
        try:
            el = page.query_selector(sel)
            if el and el.is_visible():
                return True
        except Exception:
            continue

    # 兜底: 检测 url 关键词
    url = page.url.lower()
    if "qrcode" in url or "qr" in url or "scan" in url:
        return True

    # 兜底: 检测页面文本
    text = page.inner_text("body")[:500].lower() if page.query_selector("body") else ""
    for kw in ["扫码", "扫描二维码", "请使用微信", "验证身份"]:
        if kw in text:
            return True
    return False


def _wait_qr_scan(page, timeout: int = 120) -> bool:
    """
    检测到二维码页面 → 截屏 → webhook 通知 → 等待扫码完成。

    返回 True 表示验证通过（已离开二维码页面）。
    """
    print("\n  [QR] 检测到二维码登录页")

    # 截屏
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    screenshot_path = _DATA_DIR / f"qr_verify_{ts}.png"
    try:
        page.screenshot(path=str(screenshot_path))
        print(f"  [QR] 截屏已保存: {screenshot_path.name}")
    except Exception as e:
        print(f"  [QR] 截屏失败: {e}")
        screenshot_path = None

    # 通过 herald_agent 发送通知
    try:
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
        from herald_agent import send as herald_send
        msg = (
            f"需要扫码验证登录\n"
            f"截屏: {screenshot_path.name if screenshot_path else 'N/A'}\n"
            f"请在 {timeout} 秒内扫码完成验证"
        )
        herald_send(override_status="pending", override_msg=msg)
        print(f"  [QR] 通知已发送到手机")
    except Exception as e:
        print(f"  [QR] 通知发送失败（不影响等待）: {e}")

    # 等待扫码完成：轮询检测页面是否跳转
    start = time.time()
    check_interval = 3

    # 同时监听 flag 文件（手动确认）
    if _VERIFY_FLAG.exists():
        _VERIFY_FLAG.unlink()

    print(f"  [QR] 等待扫码...（超时 {timeout}s）")
    print(f"  [QR] 扫码后自动继续，或创建文件: {_VERIFY_FLAG}")

    while time.time() - start < timeout:
        # 1. 检测 flag 文件
        if _VERIFY_FLAG.exists():
            print(f"  [QR] flag 文件检测到，继续执行")
            _VERIFY_FLAG.unlink()
            return True

        # 2. 检测页面是否跳转（URL 变化 / 二维码消失）
        try:
            if not _is_qr_page(page):
                print(f"  [QR] 页面已跳转，验证通过")
                return True
        except Exception:
            return True  # 页面已关闭/导航

        time.sleep(check_interval)

    print(f"  [QR] 超时 ({timeout}s)，跳过验证继续执行")
    return False


def _parse_num(s: str) -> float:
    """解析数字字符串（含万/亿）"""
    s = s.replace(",", "").strip()
    if not s:
        return 0
    try:
        if "万" in s:
            return float(s.replace("万", "")) * 10000
        if "亿" in s:
            return float(s.replace("亿", "")) * 100000000
        return float(s)
    except ValueError:
        return 0


# ──────────────────────────────────────────────
# 竞对/大号评论区抓取
# ──────────────────────────────────────────────

_MOCK_COMMENTS = {
    "知识分享": [
        "怎么才能记得住这么多知识点？",
        "有没有适合上班族的学习时间安排？",
        "这些方法真的有用吗？实测过没有？",
        "看完了还是不知道怎么开始",
        "能不能推荐一些入门级的书单",
        "为什么我照着做效果不明显？",
        "有没有更简单的方法？",
        "这些技巧适合零基础吗？",
        "你是怎么坚持下来的？",
        "能不能出一期慢节奏的教学视频？",
        "内容太干，能不能多加点实操案例？",
        "你用了多久才形成这个体系的？",
        "看完收藏等于学会吗？",
        "有没有交流群可以一起学习？",
    ],
    "创业": [
        "启动资金最少需要多少？",
        "失败了你怎么办？有退路吗？",
        "现在行情不好，适合创业吗？",
        "能不能带人一起做？",
        "怎么找到第一批客户？",
        "这个模式在小城市能跑通吗？",
        "投入产出比大概多少？",
        "有没有不需要太多本钱的项目推荐？",
        "你是怎么平衡主业和副业的？",
        "做了多久才开始盈利的？",
        "会不会被割韭菜？怎么分辨靠谱项目？",
        "需要注册公司吗？前期要办哪些手续？",
        "团队一开始几个人比较合适？",
        "如果第一年没盈利还要坚持吗？",
    ],
    "理财": [
        "月薪5000真的有必要理财吗？",
        "基金亏了快20%要不要割肉？",
        "银行存款和理财哪个更划算？",
        "怎么判断一个理财产品靠不靠谱？",
        "有房贷还要不要投资？",
        "黄金现在还能买吗？",
        "年轻人应该先攒钱还是先投资自己？",
        "能不能推荐几个稳健型的产品？",
        "通货膨胀这么高，存钱是不是反而亏了？",
        "可转债打新现在还值得参与吗？",
        "理财会不会越理越少？",
        "有没有适合懒人的定投方案？",
        "资产配置的比例多久调整一次？",
        "怎么防止被银行理财经理忽悠？",
    ],
    "AI": [
        "AI会不会取代我的工作？",
        "哪个AI工具最好用？免费的有推荐吗？",
        "用AI生成的内容会被平台判违规吗？",
        "提示词怎么写才能出高质量结果？",
        "AI画画能商用吗？版权怎么算？",
        "现在学AI还来得及吗？门槛高不高？",
        "有没有适合新手的AI工具推荐？",
        "AI写的文案会不会被检测出来？",
        "怎么用AI提升工作效率？有具体案例吗？",
        "AI生成的代码安全吗？能用在生产环境？",
        "国内有哪些好用的AI模型？",
        "AI视频生成现在什么水平了？",
        "会不会用AI的人替代不会用的人？",
        "AI标注的内容在抖音会不会被限流？",
    ],
}

_MOCK_COMMENT_FLAT = [
    {"text": c, "source": domain, "category": "用户疑问"}
    for domain, comments in _MOCK_COMMENTS.items()
    for c in comments
]


def _scrape_comments(cfg: dict) -> list:
    """Playwright 抓取竞对大号评论区"""
    accounts = cfg.get("competitor", {}).get("accounts", [])
    max_per = cfg.get("competitor", {}).get("max_comments_per_account", 30)
    enabled = cfg.get("competitor", {}).get("enabled", False)

    if not enabled or not accounts:
        print("[SKIP] 竞对评论抓取未启用")
        return []

    from playwright.sync_api import sync_playwright

    all_comments = []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True, args=["--no-sandbox"])
        ctx = browser.new_context(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 Chrome/124.0.0.0 Safari/537.36"
            ),
            locale="zh-CN",
        )
        for acct in accounts:
            url = acct.get("url", "")
            label = acct.get("label", "?")
            if not url:
                continue
            try:
                page = ctx.new_page()
                page.goto(url, wait_until="domcontentloaded", timeout=30000)
                page.wait_for_timeout(3000)
                # 滚动加载更多评论
                for _ in range(3):
                    page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
                    page.wait_for_timeout(2000)
                # 提取评论
                selectors = [
                    ".comment-content",
                    ".comment-item span",
                    "[data-testid='comment']",
                    ".reply-content",
                ]
                comments = []
                for sel in selectors:
                    els = page.query_selector_all(sel)
                    for el in els:
                        t = el.inner_text().strip()
                        if t and len(t) > 3 and t not in comments:
                            comments.append(t)
                    if len(comments) >= max_per:
                        break
                for c in comments[:max_per]:
                    all_comments.append({"text": c, "source": label, "category": "用户评论"})
                print(f"  [SCRAPE] {label}: {len(comments)} 条")
                page.close()
            except Exception as e:
                print(f"  [FAIL] {label}: {e}")
        ctx.close()
        browser.close()
    return all_comments


def _extract_user_pains(comments: list, max_pains: int = 50) -> list:
    """简单频率统计 + 关键词聚类 → 提取用户痛点"""
    from collections import Counter

    # 高频问题词
    pain_keywords = [
        "怎么", "如何", "能不能", "有没有", "适合", "值得",
        "安全", "靠谱", "门槛", "成本", "风险", "效果",
        "替代", "开始", "坚持", "赚钱", "亏", "买", "推荐",
    ]

    # 按关键词聚类
    clusters = Counter()
    for c in comments:
        text = c.get("text", "")
        for kw in pain_keywords:
            if kw in text:
                clusters[kw] += 1

    # 提取痛点
    pains = []
    seen = set()
    for c in comments:
        text = c.get("text", "").strip()
        if not text or text in seen:
            continue
        seen.add(text)
        # 计算 pain score = 关键词命中数
        score = sum(1 for kw in pain_keywords if kw in text)
        pains.append({
            "text": text,
            "source": c.get("source", "未知"),
            "score": score,
            "keywords": [kw for kw in pain_keywords if kw in text],
        })

    # 按 score 排序
    pains.sort(key=lambda x: x["score"], reverse=True)
    return pains[:max_pains]


def _save_user_pains(pains: list, source_count: int):
    """写入 data/user_pains.json"""
    path = _DATA_DIR / "user_pains.json"
    data = {
        "extracted_at": datetime.now().isoformat(),
        "total_comments_scraped": source_count,
        "total_pains": len(pains),
        "pains": pains,
    }
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[PAINS] {path.name} ({len(pains)} 条痛点)")
    return data

def _write_excel(data: list):
    """写入 feedback_latest.xlsx"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "发布反馈"

    # 表头
    ws.append(_COLUMNS)
    for col in range(1, len(_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")

    # 数据
    for row_data in data:
        if isinstance(row_data, dict):
            engagement = 0
            plays = row_data.get("plays", 0) or 0
            if plays > 0:
                likes = row_data.get("likes", 0) or 0
                comments = row_data.get("comments", 0) or 0
                shares = row_data.get("shares", 0) or 0
                engagement = (likes + comments + shares) / plays
            baseline = row_data.get("baseline_engagement", engagement * 1.2)

            ws.append([
                row_data.get("case_id", ""),
                row_data.get("original_text", ""),
                row_data.get("refined_text", ""),
                row_data.get("strategy", "standard"),
                plays, row_data.get("likes", 0), row_data.get("comments", 0),
                row_data.get("shares", 0),
                round(engagement, 4), round(baseline, 4),
            ])
        elif isinstance(row_data, tuple):
            # (case_id, orig, refined, strategy, plays, likes, comments, shares, baseline)
            case_id, orig, refined, strategy, plays, likes, comments, shares, baseline = row_data
            engagement = (likes + comments + shares) / plays if plays > 0 else 0
            ws.append([case_id, orig, refined, strategy, plays, likes, comments, shares,
                       round(engagement, 4), baseline])

    wb.save(str(_FEEDBACK_PATH))
    print(f"[EXCEL] {_FEEDBACK_PATH.name} ({len(data)} 条)")


# ──────────────────────────────────────────────
# 触发自进化
# ──────────────────────────────────────────────

def _trigger_evolution():
    """调用 feedback_listener.py 的自进化逻辑"""
    print("[TRIGGER] feedback_listener.py ...")
    from feedback_listener import listen as fb_listen
    fb_listen(feedback_paths=[_FEEDBACK_PATH])


# ──────────────────────────────────────────────
# 主入口
# ──────────────────────────────────────────────

def collect(mock: bool = False, scrape_comments: bool = False):
    """采集数据 → 导出 → 自进化 · 可选竞对评论抓取"""
    print("=" * 48)
    print("  规则甄查 · 甄先生 v2.0 — 反馈自动采集")
    print("=" * 48)

    # ── 竞对评论抓取 ──
    if scrape_comments:
        print("\n[COMMENTS] 竞对评论抓取模式\n")
        cfg = yaml.safe_load(_CFG_PATH.read_text(encoding="utf-8")) if _CFG_PATH.exists() else {}
        if mock:
            comments = _MOCK_COMMENT_FLAT
            print(f"  [MOCK] {len(comments)} 条模拟评论")
        else:
            comments = _scrape_comments(cfg)

        if comments:
            pains = _extract_user_pains(comments, max_pains=50)
            _save_user_pains(pains, len(comments))

            print(f"\n[PAINS] Top 3:")
            for p in pains[:3]:
                print(f"  - [{p['score']}] {p['text'][:50]}")
        else:
            print("[SKIP] 未采集到评论")

        print(f"\n[DONE] 评论采集完成")
        return

    # ── 原有反馈采集 ──
    if mock:
        if scrape_comments:
            pass  # 已在上方处理
        else:
            _mock_feedback()
        print("\n[DONE] 模拟数据已注入自进化流程")
        return

    # 真实采集
    cfg = yaml.safe_load(_CFG_PATH.read_text(encoding="utf-8")) if _CFG_PATH.exists() else {}
    fc = cfg.get("feedback_collector", {})

    if not fc.get("enabled") or not fc.get("login_url"):
        print("[HINT] 启用真实采集需在 config.yaml 配置 feedback_collector")
        print("[HINT] 使用 --mock 生成模拟数据")
        return

    data = _real_collect(cfg)
    if not data:
        return

    _write_excel(data)

    # 触发自进化
    _trigger_evolution()

    print(f"\n[DONE] 采集完成，已触发自进化")


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="反馈自动采集器 — Playwright 登录后台抓取播放数据")
    parser.add_argument("--mock", action="store_true", help="使用模拟数据（无需配置）")
    parser.add_argument("--comments", action="store_true", help="竞对评论区抓取模式")
    args = parser.parse_args()
    collect(mock=args.mock, scrape_comments=args.comments)
